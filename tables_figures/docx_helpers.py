"""Formatting helpers for the Word tables (python-docx). Three-line academic tables."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MINUS = "−"
FONT = "Times New Roman"
TEXT_W = 6.6          # inches of usable width on A4 with 0.83-inch margins

def f(x, d=2, plus=False):
    if x is None: return "–"
    s = f"{x:.{d}f}"
    if float(s) == 0: s = s.lstrip("-")          # no "-0.00"
    if s.startswith("-"): s = MINUS + s[1:]
    elif plus and float(s) > 0: s = "+" + s
    return s

def pf(p):
    if p is None: return "–"
    if p < 0.001: return "< .001"
    return f"{p:.3f}".lstrip("0") if p < 1 else "1.000"

def pp(p):
    """'p < .001' or 'p = .023' for running text."""
    return "p < .001" if p is not None and p < 0.001 else f"p = {pf(p)}"

def st(p):
    if p is None: return ""
    return "***" if p < .001 else "**" if p < .01 else "*" if p < .05 else ""

def bse(b, se, p, d=2):
    """Coefficient with stars on line 1, (SE) on line 2."""
    return (f(b, d) + st(p), f"({f(se, d)})")

def new_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    for side in ("left_margin", "right_margin"): setattr(sec, side, Inches(0.83))
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    st_ = doc.styles["Normal"]; st_.font.name = FONT; st_.font.size = Pt(11)
    st_.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st_.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(14)
    p = doc.add_paragraph(); r = p.add_run(subtitle); r.italic = True; r.font.size = Pt(10.5)
    return doc

def para(doc, text, size=11, bold=False, italic=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    return p

def heading(doc, text, size=12.5, before=10):
    p = para(doc, text, size=size, bold=True, after=4, before=before)
    p.paragraph_format.keep_with_next = True
    return p

def _border(cell, edge, sz=8, color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:color"), color)
    borders.append(el)

def _write(cell, content, size, bold=False, italic=False, align="center"):
    lines = content if isinstance(content, (list, tuple)) else [content]
    cell.paragraphs[0].text = ""
    for k, line in enumerate(lines):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        r = p.add_run(str(line)); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.name = FONT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, number, title, header, rows, widths, note=None, size=9, interpretation=None):
    """rows: list of lists; a row given as ("PANEL", "text") becomes a full-width italic panel line."""
    assert abs(sum(widths) - TEXT_W) < 0.05, f"{number}: widths sum {sum(widths)}"
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{number}. "); r.bold = True; r.font.size = Pt(11)
    r = p.add_run(title); r.italic = True; r.font.size = Pt(11)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, w in enumerate(widths): t.columns[j].width = Inches(w)      # grid widths (LibreOffice/Google Docs)
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))): gc.set(qn("w:w"), str(int(widths[j] * 1440)))
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.width = Inches(widths[j])
        _write(c, h, size, bold=True, align="left" if j == 0 else "center")
        _border(c, "top", 10); _border(c, "bottom", 6)
    for i, row in enumerate(rows, start=1):
        cells = t.rows[i].cells
        for j in range(len(header)): cells[j].width = Inches(widths[j])
        if isinstance(row, tuple) and row[0] == "PANEL":
            m = cells[0].merge(cells[-1]); _write(m, row[1], size, italic=True, align="left")
            if i > 1: _border(m, "top", 4, "808080")
            continue
        for j, val in enumerate(row):
            _write(cells[j], val, size, align="left" if j == 0 else "center")
    for c in t.rows[-1].cells: _border(c, "bottom", 10)
    # keep each table on one page where possible
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
    if note:
        q = para(doc, "", size=8.5, after=4, before=3)
        r = q.add_run("Note: "); r.italic = True; r.font.size = Pt(8.5)
        r = q.add_run(note); r.font.size = Pt(8.5)
    if interpretation:
        interpret(doc, interpretation)
    return t

def interpret(doc, text):
    q = doc.add_paragraph(); q.paragraph_format.space_before = Pt(4); q.paragraph_format.space_after = Pt(8)
    r = q.add_run("Interpretation. "); r.bold = True; r.font.size = Pt(10.5)
    r = q.add_run(text); r.font.size = Pt(10.5)

def figure(doc, number, title, path, note=None, interpretation=None, width=6.3):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{number}. "); r.bold = True; r.font.size = Pt(11)
    r = p.add_run(title); r.italic = True; r.font.size = Pt(11)
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    pic.add_run().add_picture(str(path), width=Inches(width))
    if note:
        q = para(doc, "", size=8.5, after=4)
        r = q.add_run("Note: "); r.italic = True; r.font.size = Pt(8.5)
        r = q.add_run(note); r.font.size = Pt(8.5)
    if interpretation: interpret(doc, interpretation)
